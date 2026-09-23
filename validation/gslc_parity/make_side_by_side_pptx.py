"""One-slide deck: classical vs GSLC (sign-corrected) wrapped phase at full resolution, side by side.
Images: <figures>/img/full/{classical,GSLC_after}.jpg (from plot_ifg_compare.py 2 8); figures live outside the repo.
  make_side_by_side_pptx.py [<figures_dir> [<out.pptx>]]   (default figures dir E:/ESA/gslc-parity-figures)
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIG = Path(r"E:\ESA\gslc-parity-figures")
INK, MUTED, BG, ACC = RGBColor(0x1A, 0x22, 0x30), RGBColor(0x5B, 0x66, 0x75), RGBColor(0xF3, 0xF5, 0xF7), RGBColor(0x0F, 0x76, 0x6E)


def text(slide, x, y, w, h, s, size, bold=False, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb, r.font.name = s, Pt(size), bold, color, "Arial"
    return tb


def main(fig, out):
    IMG = Path(fig) / "img" / "full"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    text(s, 0.5, 0.25, 12.3, 0.6, "Classical vs GSLC: same fringes at full resolution", 28, True)
    w = 6.1
    h = w * 2256 / 2958
    for x, name, label, col in ((0.5, "classical", "Classical (Back-Geocoding + ESD)", INK),
                                (6.73, "GSLC_after", "GSLC, corrected sign (new default)", ACC)):
        text(s, x, 1.0, w, 0.4, label, 18, True, col)
        s.shapes.add_picture(str(IMG / f"{name}.jpg"), Inches(x), Inches(1.45), Inches(w), Inches(h))
    text(s, 0.5, 1.45 + h + 0.1, 12.3, 0.8,
         "Venezuela S1A x S1C, IW3 VV, bursts 4-6 top to bottom, ETAD on, ramp off. Wrapped phase (-pi..pi) on 2 x 8 "
         "radar cells of the classical burst geometry (1 cell = 1 image pixel, 2958 x 2256). GSLC binned into the cells "
         "by recorded source position (no resampling). Zoom in the file to inspect: images are embedded at full size.",
         12, False, MUTED)
    prs.save(out)
    print(out)


if __name__ == "__main__":
    a = sys.argv[1:]
    fig = a[0] if a else FIG
    main(fig, a[1] if len(a) > 1 else str(Path(fig) / "GSLC_vs_classical_fullres.pptx"))
